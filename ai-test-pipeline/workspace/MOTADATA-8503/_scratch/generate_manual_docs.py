# Generate downloadable manual test docs (Excel + PDF) from manual-cases.yaml.
# Includes Jira links, tester-friendly detail, and an Impact/Affected section
# (only impact that is 100% known from the ticket: F1/F2 backward-compat + A3 removed flows).
import yaml, io, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

D = os.path.dirname(__file__)
JIRA = "http://172.16.10.6:8080/browse/MOTADATA-8503"
cases = yaml.safe_load(open(os.path.join(D, "manual-cases.yaml"), encoding="utf-8"))

# Impact / affected scenarios that are 100% stated in the ticket (no speculation):
IMPACT = {
  "F1": "Existing Threshold policies (Metric/APM/RUM/NetRoute) must keep working unchanged — regression risk from unified UI.",
  "F2": "Existing Anomaly/Forecast policies must still open & function in the unified screen.",
  "A3": "Standalone Anomaly/Forecast creation flows are removed — any saved links/bookmarks/docs pointing to old flows are affected.",
}
def impact_for(tr):
    out = [v for k, v in IMPACT.items() if k in (tr or [])]
    return out

def steps_text(c):
    out = []
    for s in c.get("steps", []):
        if isinstance(s, dict):
            for k, v in s.items():
                out.append(f"{k}: {v}")
        else:
            out.append(str(s))
    return out

def expected_text(c):
    out = []
    for e in c.get("expected", []):
        out.append(", ".join(f"{k}={v}" for k, v in e.items()) if isinstance(e, dict) else str(e))
    return out

# ---------------- EXCEL ----------------
wb = Workbook(); ws = wb.active; ws.title = "MOTADATA-8503 Manual Tests"
hdr = ["Test ID","Title","Jira","Module","Risk","Traces To (AC)","Preconditions",
       "Steps","Expected Result","Edge Cases","Impact / Affected (known)"]
head_fill = PatternFill("solid", fgColor="2B394F"); head_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="CCCCCC"); border = Border(thin,thin,thin,thin)
ws.append(hdr)
for i,h in enumerate(hdr,1):
    cell = ws.cell(1,i); cell.fill=head_fill; cell.font=head_font
    cell.alignment=Alignment(vertical="top", wrap_text=True); cell.border=border
for c in cases:
    row = [c["id"], c["title"], JIRA, c.get("module",""), c.get("risk",""),
           ", ".join(c.get("traces_to",[])), "\n".join(c.get("preconditions",[])),
           "\n".join(steps_text(c)), "\n".join(expected_text(c)),
           "\n".join(c.get("edge_cases",[])), "\n".join(impact_for(c.get("traces_to",[]))) or "—"]
    ws.append(row)
    r = ws.max_row
    ws.cell(r,3).hyperlink = JIRA; ws.cell(r,3).font = Font(color="0563C1", underline="single")
    for ci in range(1,len(hdr)+1):
        cc=ws.cell(r,ci); cc.alignment=Alignment(vertical="top", wrap_text=True); cc.border=border
widths=[12,34,16,12,8,16,30,52,34,30,40]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes="A2"
xlsx=os.path.join(D,"MOTADATA-8503_ManualTests.xlsx"); wb.save(xlsx)

# ---------------- PDF ----------------
styles=getSampleStyleSheet()
h1=ParagraphStyle("h1",parent=styles["Heading1"],fontSize=15,textColor=colors.HexColor("#2B394F"))
h2=ParagraphStyle("h2",parent=styles["Heading2"],fontSize=11,textColor=colors.HexColor("#2B394F"),spaceBefore=8)
body=ParagraphStyle("body",parent=styles["BodyText"],fontSize=8.5,leading=12)
small=ParagraphStyle("small",parent=styles["BodyText"],fontSize=8,textColor=colors.grey)
pdf=os.path.join(D,"MOTADATA-8503_ManualTests.pdf")
doc=SimpleDocTemplate(pdf,pagesize=A4,topMargin=14*mm,bottomMargin=12*mm,leftMargin=14*mm,rightMargin=14*mm)
el=[]
el.append(Paragraph("MOTADATA-8503 — Manual Test Cases",h1))
el.append(Paragraph(f'Unified Create Policy UI (Metric/APM/RUM/NetRoute). Jira: <a href="{JIRA}">{JIRA}</a>',small))
el.append(Spacer(1,6))
for c in cases:
    el.append(Paragraph(f'{c["id"]} — {c["title"]}',h2))
    el.append(Paragraph(f'<b>Module:</b> {c.get("module","")} &nbsp; <b>Risk:</b> {c.get("risk","")} &nbsp; '
                        f'<b>Traces to:</b> {", ".join(c.get("traces_to",[]))} &nbsp; '
                        f'<b>Jira:</b> <a href="{JIRA}">{c["id"]}</a>',small))
    if c.get("preconditions"):
        el.append(Paragraph("<b>Preconditions:</b> "+"; ".join(c["preconditions"]),body))
    el.append(Paragraph("<b>Steps:</b>",body))
    for i,s in enumerate(steps_text(c),1): el.append(Paragraph(f"{i}. {s}",body))
    if expected_text(c): el.append(Paragraph("<b>Expected:</b> "+"; ".join(expected_text(c)),body))
    if c.get("edge_cases"): el.append(Paragraph("<b>Edge cases:</b> "+"; ".join(c["edge_cases"]),body))
    imp=impact_for(c.get("traces_to",[]))
    if imp:
        el.append(Paragraph("<b>Impact / Affected (known):</b>",body))
        for x in imp: el.append(Paragraph("• "+x,body))
    el.append(Spacer(1,8))
doc.build(el)
print("WROTE", xlsx); print("WROTE", pdf)
